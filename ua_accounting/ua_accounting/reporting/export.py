"""Excel export of a table model (the same model the browser draws)."""

import re
from io import BytesIO

from frappe.utils import get_url_to_form
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

MONEY = "#,##0.00;-#,##0.00;;"
QTY = "#,##0.000;-#,##0.000;;"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E8E8E8")
SUM_FILL = PatternFill("solid", fgColor="F3F3F3")
WIDTH = {"label": 46, "money": 16, "qty": 13, "text": 30, "date": 12, "doc": 30, "lines": 28}


def to_xlsx(table):
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", table["title"] or "Report")[:31]

    r = 1
    ws.cell(r, 1, table["title"]).font = Font(bold=True, size=14)
    r += 1
    for line in table["lines"]:
        ws.cell(r, 1, line)
        r += 1
    r += 1

    # header: place cells on a grid, a cell may span rows and columns
    busy = set()
    for hrow in table["header"]:
        col = 1
        for cell in hrow:
            while (r, col) in busy:
                col += 1
            rs, cs = cell.get("rowspan", 1), cell.get("colspan", 1)
            for dr in range(rs):
                for dc in range(cs):
                    busy.add((r + dr, col + dc))
                    ws.cell(r + dr, col + dc).border = BORDER
                    ws.cell(r + dr, col + dc).fill = HEAD_FILL
            c = ws.cell(r, col, cell["label"])
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if rs > 1 or cs > 1:
                ws.merge_cells(start_row=r, start_column=col, end_row=r + rs - 1, end_column=col + cs - 1)
            col += cs
        r += 1
    ws.freeze_panes = ws.cell(r, 1)

    for row in table["rows"]:
        for i, col in enumerate(table["columns"], start=1):
            value = row["cells"].get(col["key"])
            kind = col["kind"]
            if isinstance(value, (list, tuple)):
                value = "\n".join(str(v) for v in value)
            c = ws.cell(r, i, value if value not in (None, "") else None)
            c.border = BORDER
            c.font = Font(bold=bool(row.get("bold")))
            if row["kind"] in ("total", "open", "close"):
                c.fill = SUM_FILL
            if kind == "money":
                c.number_format = MONEY
                c.alignment = Alignment(horizontal="right")
            elif kind == "qty":
                c.number_format = QTY
                c.alignment = Alignment(horizontal="right")
            elif kind == "label":
                c.alignment = Alignment(indent=row.get("level", 0), vertical="top")
            elif kind == "lines":
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(horizontal=col.get("align", "left"), vertical="top")
            link = row.get("links", {}).get(col["key"])
            if link:
                c.hyperlink = get_url_to_form(link["doctype"], link["name"])
                c.font = Font(bold=bool(row.get("bold")), color="0563C1", underline="single")
        r += 1

    if table.get("message"):
        r += 1
        ws.cell(r, 1, table["message"])

    for i, col in enumerate(table["columns"], start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(col["kind"], 16)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
