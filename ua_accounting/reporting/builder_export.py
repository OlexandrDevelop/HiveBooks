"""Excel and PDF files of the universal report result."""

from io import BytesIO

from frappe import _
from frappe.utils import escape_html, flt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ua_accounting.accounting.utils import uk_money

MONEY = "#,##0.00"
NUMBER_TYPES = ("Currency", "Float", "Int", "Percent")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E8E8E8")
SUM_FILL = PatternFill("solid", fgColor="F3F3F3")


def cell_value(value, fieldtype):
    if value is None or value == "":
        return None
    if fieldtype in NUMBER_TYPES:
        return flt(value)
    return str(value)


def to_xlsx(title, result):
    wb = Workbook()
    ws = wb.active
    ws.title = "Звіт"
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    r = 3
    columns = result["columns"]
    for i, col in enumerate(columns, start=1):
        c = ws.cell(r, i, col["label"])
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(r + 1, 1)
    r += 1
    for row in result["rows"]:
        level = row.get("_level")
        bold = bool(level)
        for i, col in enumerate(columns, start=1):
            c = ws.cell(r, i, cell_value(row.get(col["fieldname"]), col["fieldtype"]))
            c.border = BORDER
            style = (row.get("_format") or {}).get(col["fieldname"]) or {}
            c.font = Font(bold=bold or bool(style.get("bold")), color=(style.get("color") or "#000000").lstrip("#"))
            if col["fieldtype"] in NUMBER_TYPES:
                c.number_format = MONEY if col["fieldtype"] == "Currency" else "General"
                c.alignment = Alignment(horizontal="right")
            elif level and i == level:
                c.alignment = Alignment(indent=level - 1)
        r += 1
    totals = result["totals"]
    label_done = False
    for i, col in enumerate(columns, start=1):
        value = totals.get(col["fieldname"])
        text = None
        if value is None and not label_done:
            text, label_done = _("Разом"), True
        c = ws.cell(r, i, flt(value) if value is not None else text)
        c.font = Font(bold=True)
        c.fill = SUM_FILL
        c.border = BORDER
        if value is not None:
            c.number_format = MONEY if col["fieldtype"] == "Currency" else "General"
            c.alignment = Alignment(horizontal="right")
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def to_html(title, result):
    def show(value, fieldtype):
        if value is None or value == "":
            return ""
        if fieldtype == "Currency":
            return uk_money(value)
        return escape_html(str(value))

    head = "".join(f"<th>{escape_html(c['label'])}</th>" for c in result["columns"])
    body = []
    for row in result["rows"]:
        level = row.get("_level")
        cells = []
        for col in result["columns"]:
            style = (row.get("_format") or {}).get(col["fieldname"]) or {}
            css = ("font-weight:bold;" if level or style.get("bold") else "") + (f"color:{escape_html(style['color'])};" if style.get("color") else "")
            right = "text-align:right;" if col["fieldtype"] in NUMBER_TYPES else ""
            cells.append(f'<td style="{css}{right}">{show(row.get(col["fieldname"]), col["fieldtype"])}</td>')
        body.append("<tr>" + "".join(cells) + "</tr>")
    foot, done = [], False
    for col in result["columns"]:
        value = result["totals"].get(col["fieldname"])
        if value is None:
            foot.append(f"<td><b>{'' if done else escape_html(_('Разом'))}</b></td>")
            done = True
        else:
            foot.append(f'<td style="text-align:right"><b>{show(value, col["fieldtype"])}</b></td>')
    return (
        '<html><head><meta charset="utf-8"><style>'
        "body{font-family:sans-serif;font-size:11px}table{border-collapse:collapse;width:100%}"
        "th,td{border:1px solid #999;padding:3px 6px}th{background:#e8e8e8}</style></head><body>"
        f"<h3>{escape_html(title)}</h3><table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody>"
        f"<tfoot><tr>{''.join(foot)}</tr></tfoot></table></body></html>"
    )
