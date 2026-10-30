"""Excel templates: headers in Ukrainian, one example row."""

import base64
import io
from datetime import date, datetime

import frappe
from frappe import _

from ua_accounting.exchange.common import ExchangeError, flag
from ua_accounting.exchange.onec import parse_date

HEADERS = {
    "counterparties": ["Найменування", "ЄДРПОУ", "ІПН", "Платник ПДВ", "Покупець", "Постачальник"],
    "items": ["Найменування", "Артикул", "Одиниця", "Ставка ПДВ", "Вид"],
    "warehouses": ["Найменування"],
    "balances": ["Дата", "Рахунок", "Контрагент", "Договір", "Склад", "Номенклатура", "Рахунок IBAN", "Кількість", "Сума", "Сторона"],
}
EXAMPLES = {
    "counterparties": ["ТОВ «Приклад»", "12345678", "", "Так", "Так", "Ні"],
    "items": ["Борошно пшеничне", "BR-001", "кг", "20", "Товар"],
    "warehouses": ["Основний склад"],
    "balances": ["2025-12-31", "311", "", "", "", "", "UA043052990000026001000000001", "", "1000.00", "Дт"],
}
TITLES = {"counterparties": "Контрагенти", "items": "Номенклатура", "warehouses": "Склади", "balances": "Залишки"}


def template(kind):
    import openpyxl
    from openpyxl.styles import Font

    if kind not in HEADERS:
        frappe.throw(_("Невідомий шаблон «{0}»").format(kind))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TITLES[kind]
    ws.append(HEADERS[kind])
    ws.append(EXAMPLES[kind])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, title in enumerate(HEADERS[kind], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(16, len(title) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return {"filename": f"shablon_{kind}.xlsx", "content": base64.b64encode(buf.getvalue()).decode("ascii")}


def text(value):
    """A cell as text: numbers lose the `.0` of Excel floats."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else ""
    if isinstance(value, float) and value == int(value):
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def read(data, kind):
    """Rows of the first sheet as dicts keyed by header; `row` is the row number in the sheet."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ExchangeError(_("Файл не є книгою Excel (.xlsx)"))
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [text(c) for c in next(rows, [])]
    expected = HEADERS[kind]
    if header[: len(expected)] != expected:
        raise ExchangeError(_("Заголовки першого рядка не збігаються з шаблоном. Потрібні: {0}").format("; ".join(expected)))
    out = []
    for n, row in enumerate(rows, start=2):
        cells = [text(c) for c in row[: len(expected)]]
        if not any(cells):
            continue
        cells += [""] * (len(expected) - len(cells))
        out.append(dict(zip(expected, cells), row=n))
    return out


def zfill_code(value):
    """ЄДРПОУ typed as a number loses its leading zeros."""
    return value.zfill(8) if value.isdigit() and 0 < len(value) < 8 else value


def counterparties(rows):
    return [{"row": r["row"], "name": r["Найменування"], "edrpou": zfill_code(r["ЄДРПОУ"]), "ipn": r["ІПН"],
             "vat_payer": flag(r["Платник ПДВ"]), "is_customer": flag(r["Покупець"]), "is_supplier": flag(r["Постачальник"])}
            for r in rows]


def items(rows):
    out = []
    for r in rows:
        rate = r["Ставка ПДВ"].replace("%", "").strip()
        out.append({"row": r["row"], "name": r["Найменування"], "article": r["Артикул"], "uom": r["Одиниця"],
                    "vat_rate": rate, "item_type": r["Вид"]})
    return out


def warehouses(rows):
    return [{"row": r["row"], "name": r["Найменування"]} for r in rows]


def balances(rows):
    """(date, rows) of the balances sheet; all rows must have one date."""
    dates = {parse_date(r["Дата"]) for r in rows if r["Дата"]}
    if None in dates:
        raise ExchangeError(_("Дату в колонці «Дата» не розпізнано. Вкажіть її як ДД.ММ.РРРР"))
    if len(dates) > 1:
        raise ExchangeError(_("Усі рядки залишків мають бути на одну дату"))
    return (dates.pop() if dates else None), [{
        "row": r["row"], "account": r["Рахунок"], "counterparty": zfill_code(r["Контрагент"]), "contract": r["Договір"],
        "warehouse": r["Склад"], "item": r["Номенклатура"], "iban": r["Рахунок IBAN"], "quantity": r["Кількість"],
        "amount": r["Сума"], "side": r["Сторона"]} for r in rows]
